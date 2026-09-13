"""
modelos/registro_acceso.py
----------------------------
MÓDULO: Registro de Accesos / Reportes
Consultas SQL sobre la tabla `registros_acceso` — cada fila es un
intento de marcación de huella en el Kiosco (permitido o denegado),
asociado a un funcionario.

Este modelo lo usan dos pantallas: registro.html (vista del momento
actual) y reportes.html (vista histórica con más filtros y
estadísticas) — ambas reusan listar_registros() con distintos
filtros, para no duplicar la misma consulta en dos archivos.
"""

from database.conexion import obtener_conexion


def listar_registros(fecha_desde=None, fecha_hasta=None, estado=None, categoria=None, texto_busqueda=None):
    """
    Devuelve los registros de acceso que cumplen los filtros dados,
    del más reciente al más antiguo, con nombre/documento/categoría
    del funcionario ya incluidos (LEFT JOIN contra funcionarios).

    Todos los filtros son opcionales — si no se pasa ninguno, trae
    todo. 'categoria' filtra por el rol real del funcionario
    (Funcionario/Visitante/Vigilancia/Administrativo). El filtro de
    texto busca coincidencia parcial en nombre, apellido o documento.
    """
    conexion = obtener_conexion()
    cursor = conexion.cursor(dictionary=True)

    condiciones = []
    valores = []

    if fecha_desde:
        condiciones.append("r.fecha >= %s")
        valores.append(fecha_desde)
    if fecha_hasta:
        condiciones.append("r.fecha <= %s")
        valores.append(fecha_hasta)
    if estado:
        condiciones.append("r.estado = %s")
        valores.append(estado)
    if categoria:
        condiciones.append("f.categoria = %s")
        valores.append(categoria)
    if texto_busqueda:
        condiciones.append("(f.nombres LIKE %s OR f.apellidos LIKE %s OR f.id_funcionario LIKE %s)")
        patron = f"%{texto_busqueda}%"
        valores.extend([patron, patron, patron])

    where_sql = ("WHERE " + " AND ".join(condiciones)) if condiciones else ""

    consulta = f"""
        SELECT
            r.id_registro, r.fecha, r.hora, r.tipo_acceso,
            r.estado, r.codigo_registro,
            f.id_funcionario, f.nombres, f.apellidos, f.categoria
        FROM registros_acceso r
        LEFT JOIN funcionarios f ON r.id_funcionario = f.id_funcionario
        {where_sql}
        ORDER BY r.fecha DESC, r.hora DESC
    """

    cursor.execute(consulta, valores)
    resultados = cursor.fetchall()

    cursor.close()
    conexion.close()

    for fila in resultados:
        if fila.get("fecha"):
            fila["fecha"] = fila["fecha"].strftime("%d/%m/%Y")
        if fila.get("hora"):
            fila["hora"] = _formatear_hora(fila["hora"])

    return resultados


def _formatear_hora(valor_hora):
    """
    mysql-connector entrega una columna TIME como timedelta, no como
    hora legible. Esta función la convierte a texto tipo '06:48 a.m.'
    """
    total_segundos = int(valor_hora.total_seconds())
    horas = (total_segundos // 3600) % 24
    minutos = (total_segundos % 3600) // 60

    sufijo = "a.m." if horas < 12 else "p.m."
    horas_12 = horas % 12 or 12

    return f"{horas_12:02d}:{minutos:02d} {sufijo}"


def obtener_estadisticas_reportes(fecha_desde=None, fecha_hasta=None, categoria=None, texto_busqueda=None):
    """
    Devuelve los 4 totales de las tarjetas de Reportes (ingreso, salida,
    permitidos, denegados) para el periodo dado, y el % de cambio contra
    el periodo inmediatamente anterior de la misma duración (solo si se
    dieron ambas fechas; si no, delta_* queda en None).
    """
    conexion = obtener_conexion()
    cursor = conexion.cursor(dictionary=True)

    def _contar(f_desde, f_hasta):
        condiciones = []
        valores = []
        if f_desde:
            condiciones.append("r.fecha >= %s")
            valores.append(f_desde)
        if f_hasta:
            condiciones.append("r.fecha <= %s")
            valores.append(f_hasta)
        if categoria:
            condiciones.append("f.categoria = %s")
            valores.append(categoria)
        if texto_busqueda:
            condiciones.append("(f.nombres LIKE %s OR f.apellidos LIKE %s OR f.id_funcionario LIKE %s)")
            patron = f"%{texto_busqueda}%"
            valores.extend([patron, patron, patron])

        where_sql = ("WHERE " + " AND ".join(condiciones)) if condiciones else ""

        consulta = f"""
            SELECT
                SUM(CASE WHEN r.tipo_acceso = 'Ingreso' THEN 1 ELSE 0 END) AS total_ingreso,
                SUM(CASE WHEN r.tipo_acceso = 'Salida' THEN 1 ELSE 0 END) AS total_salida,
                SUM(CASE WHEN r.estado = 'Permitido' THEN 1 ELSE 0 END) AS permitidos,
                SUM(CASE WHEN r.estado = 'Denegado' THEN 1 ELSE 0 END) AS denegados
            FROM registros_acceso r
            LEFT JOIN funcionarios f ON r.id_funcionario = f.id_funcionario
            {where_sql}
        """
        cursor.execute(consulta, valores)
        fila = cursor.fetchone()
        return {clave: int(valor or 0) for clave, valor in fila.items()}
    actual = _contar(fecha_desde, fecha_hasta)

    anterior = None
    if fecha_desde and fecha_hasta:
        from datetime import date, datetime, timedelta

        d_desde = fecha_desde if isinstance(fecha_desde, date) else datetime.strptime(fecha_desde, "%Y-%m-%d").date()
        d_hasta = fecha_hasta if isinstance(fecha_hasta, date) else datetime.strptime(fecha_hasta, "%Y-%m-%d").date()
        dias = (d_hasta - d_desde).days + 1

        fecha_hasta_anterior = d_desde - timedelta(days=1)
        fecha_desde_anterior = fecha_hasta_anterior - timedelta(days=dias - 1)
        anterior = _contar(fecha_desde_anterior, fecha_hasta_anterior)

    cursor.close()
    conexion.close()

    def _delta(valor_actual, valor_anterior):
        if anterior is None or not valor_anterior:
            return None
        return round(((valor_actual - valor_anterior) / valor_anterior) * 100, 1)

    return {
        "total_ingreso": actual["total_ingreso"],
        "total_salida": actual["total_salida"],
        "permitidos": actual["permitidos"],
        "denegados": actual["denegados"],
        "delta_ingreso": _delta(actual["total_ingreso"], anterior["total_ingreso"]) if anterior else None,
        "delta_salida": _delta(actual["total_salida"], anterior["total_salida"]) if anterior else None,
        "delta_permitidos": _delta(actual["permitidos"], anterior["permitidos"]) if anterior else None,
        "delta_denegados": _delta(actual["denegados"], anterior["denegados"]) if anterior else None,
    }


def obtener_tendencia_accesos(fecha_desde=None, fecha_hasta=None, categoria=None, texto_busqueda=None):
    """
    Devuelve, agrupado por día, el número de accesos Permitidos y
    Denegados dentro del periodo — esto alimenta la gráfica de barras.
    """
    conexion = obtener_conexion()
    cursor = conexion.cursor(dictionary=True)

    condiciones = []
    valores = []
    if fecha_desde:
        condiciones.append("r.fecha >= %s")
        valores.append(fecha_desde)
    if fecha_hasta:
        condiciones.append("r.fecha <= %s")
        valores.append(fecha_hasta)
    if categoria:
        condiciones.append("f.categoria = %s")
        valores.append(categoria)
    if texto_busqueda:
        condiciones.append("(f.nombres LIKE %s OR f.apellidos LIKE %s OR f.id_funcionario LIKE %s)")
        patron = f"%{texto_busqueda}%"
        valores.extend([patron, patron, patron])

    where_sql = ("WHERE " + " AND ".join(condiciones)) if condiciones else ""

    consulta = f"""
        SELECT
            r.fecha,
            SUM(CASE WHEN r.estado = 'Permitido' THEN 1 ELSE 0 END) AS permitidos,
            SUM(CASE WHEN r.estado = 'Denegado' THEN 1 ELSE 0 END) AS denegados
        FROM registros_acceso r
        LEFT JOIN funcionarios f ON r.id_funcionario = f.id_funcionario
        {where_sql}
        GROUP BY r.fecha
        ORDER BY r.fecha ASC
    """
    cursor.execute(consulta, valores)
    resultados = cursor.fetchall()
    cursor.close()
    conexion.close()

    for fila in resultados:
        if fila.get("fecha"):
            fila["fecha"] = fila["fecha"].strftime("%d/%m")
        fila["permitidos"] = int(fila["permitidos"] or 0)
        fila["denegados"] = int(fila["denegados"] or 0)

    return resultados

  # Excel
def generar_excel_reportes(ruta_destino, fecha_desde=None, fecha_hasta=None, estado=None, categoria=None, texto_busqueda=None):
    """
    Genera un archivo .xlsx en ruta_destino con los registros de acceso
    que cumplen los filtros dados (los mismos que usa la tabla de
    Reportes). Reutiliza listar_registros() para no duplicar la consulta.
    Devuelve cuántos registros se exportaron.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    registros = listar_registros(
        fecha_desde=fecha_desde,
        fecha_hasta=fecha_hasta,
        estado=estado,
        categoria=categoria,
        texto_busqueda=texto_busqueda,
    )

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Reportes de Acceso"

    encabezados = ["ID", "Fecha", "Hora", "Usuario", "Documento", "Rol", "Tipo de Acceso", "Estado", "Código de registro"]
    hoja.append(encabezados)

    fuente_encabezado = Font(bold=True, color="FFFFFF")
    relleno_encabezado = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    for columna in range(1, len(encabezados) + 1):
        celda = hoja.cell(row=1, column=columna)
        celda.font = fuente_encabezado
        celda.fill = relleno_encabezado
        celda.alignment = Alignment(horizontal="center")

    relleno_permitido = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
    fuente_permitido = Font(color="065F46", bold=True)
    relleno_denegado = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    fuente_denegado = Font(color="991B1B", bold=True)

    fila_actual = 2
    for r in registros:
        nombre_completo = f"{r['nombres']} {r['apellidos']}" if r.get("nombres") else "Desconocido"
        estado_normalizado = str(r["estado"]).strip().capitalize()

        hoja.append([
            r["id_registro"],
            r["fecha"],
            r["hora"],
            nombre_completo,
            r.get("id_funcionario") or "-",
            r.get("categoria") or "-",
            r["tipo_acceso"],
            estado_normalizado,
            r["codigo_registro"],
        ])

        celda_estado = hoja.cell(row=fila_actual, column=8)
        celda_estado.alignment = Alignment(horizontal="center")
        if estado_normalizado == "Permitido":
            celda_estado.fill = relleno_permitido
            celda_estado.font = fuente_permitido
        else:
            celda_estado.fill = relleno_denegado
            celda_estado.font = fuente_denegado

        fila_actual += 1

    anchos = [6, 12, 12, 28, 14, 16, 14, 12, 16]
    for indice, ancho in enumerate(anchos, start=1):
        letra_columna = get_column_letter(indice)
        hoja.column_dimensions[letra_columna].width = ancho

    hoja.freeze_panes = "A2"

    ultima_fila = fila_actual - 1
    if ultima_fila >= 2:
        hoja.auto_filter.ref = f"A1:{get_column_letter(len(encabezados))}{ultima_fila}"

    libro.save(ruta_destino)
    return len(registros)